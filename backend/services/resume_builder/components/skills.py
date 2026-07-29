from docx.shared import Pt

from ..helpers import add_heading, set_font


def render_skills(container, data):

    add_heading(container, "Technical Skills")

    skill_categories = data.get("skill_categories", [])

    if skill_categories:

        for category in skill_categories:

            category_name = (
                category.get("category", "")
                or "Other Skills"
            )

            skills = category.get("skills", [])

            if not skills:
                continue

            p = container.add_paragraph()

            p.space_before = Pt(1)
            p.space_after = Pt(1)

            title = p.add_run(f"{category_name}: ")
            set_font(title, size=10, bold=True)

            value = p.add_run(", ".join(skills))
            set_font(value, size=10)

    else:

        skills = data.get("skills", [])

        if isinstance(skills, list) and skills:

            p = container.add_paragraph()

            p.space_before = Pt(1)
            p.space_after = Pt(1)

            run = p.add_run(", ".join(skills))
            set_font(run, size=10)

        elif isinstance(skills, str) and skills.strip():

            p = container.add_paragraph()

            p.space_before = Pt(1)
            p.space_after = Pt(1)

            run = p.add_run(skills.strip())
            set_font(run, size=10)