from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt

from ..helpers import add_heading, add_bullet, set_font


def render_experience(container, data):

    experiences = data.get("experience", [])

    if not experiences:
        return

    add_heading(container, "Professional Experience")

    for exp in experiences:

        company = (
            exp.get("company")
            or exp.get("organization")
            or exp.get("employer")
            or ""
        )

        designation = (
            exp.get("designation")
            or exp.get("title")
            or exp.get("role")
            or ""
        )

        location = exp.get("location", "")

        start = exp.get("start_date", "")
        end = exp.get("end_date", "")

        date_text = ""

        if start and end:
            date_text = f"{start} - {end}"
        elif start:
            date_text = start
        elif end:
            date_text = end

        # -------------------------------------------------
        # Company + Dates
        # -------------------------------------------------

        p = container.add_paragraph()

        p.space_before = Pt(8)
        p.space_after = Pt(1)

        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(
            Inches(6.2),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.SPACES
        )

        left = p.add_run(company)
        set_font(left, size=11, bold=True)

        if date_text:
            right = p.add_run("\t" + date_text)
            set_font(right, size=10)

        # -------------------------------------------------
        # Designation + Location
        # -------------------------------------------------

        if designation or location:

            p = container.add_paragraph()

            p.space_before = Pt(0)
            p.space_after = Pt(2)

            text = designation

            if location:
                if text:
                    text += f" | {location}"
                else:
                    text = location

            run = p.add_run(text)

            set_font(run, size=10)

        # -------------------------------------------------
        # Responsibilities
        # -------------------------------------------------

        descriptions = exp.get("description", [])

        for item in descriptions:
            add_bullet(container, item)

        container.add_paragraph()