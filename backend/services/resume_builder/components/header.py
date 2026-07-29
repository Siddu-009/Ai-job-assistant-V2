from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from ..helpers import set_font
from ..styles import (
    PRIMARY,
    TITLE_SIZE,
    SUBTITLE_SIZE,
)


def render_header(doc, data):

    personal = data.get("personal_info", {})

    name = (
        personal.get("name")
        or data.get("name")
        or ""
    )

    title = (
        personal.get("title")
        or data.get("title")
        or ""
    )

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(name)

    set_font(
        run,
        size=TITLE_SIZE,
        bold=True,
        color=PRIMARY
    )

    if title:

        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(title)

        set_font(
            run,
            size=SUBTITLE_SIZE
        )