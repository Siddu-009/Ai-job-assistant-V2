from docx.shared import Pt

from ..helpers import add_heading, add_bullet, set_font


def render_projects(container, data):

    projects = (
        data.get("projects")
    )
    
    if not projects:
        return

    add_heading(container, "Projects")

    for project in projects:

        name = (
            project.get("name")
            or project.get("title")
            or ""
        )

        role = project.get("role", "")

        # ----------------------------------------
        # Project Name
        # ----------------------------------------

        p = container.add_paragraph()

        p.space_before = Pt(8)
        p.space_after = Pt(1)

        run = p.add_run(name)
        set_font(run, size=11, bold=True)

        if role:

            role_run = p.add_run(f"  |  {role}")
            set_font(role_run, size=10)

        # ----------------------------------------
        # Technologies
        # ----------------------------------------

        tech_stack = project.get("tech_stack", [])

        if tech_stack:

            p = container.add_paragraph()

            p.space_before = Pt(0)
            p.space_after = Pt(2)

            title = p.add_run("Technologies: ")
            set_font(title, size=10, bold=True)

            value = p.add_run(", ".join(tech_stack))
            set_font(value, size=10)

        # ----------------------------------------
        # Description
        # ----------------------------------------

        for item in project.get("description", []):

            add_bullet(container, item)

        container.add_paragraph()