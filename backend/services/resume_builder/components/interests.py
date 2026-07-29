from docx.shared import Pt

from ..helpers import add_heading, set_font


def render_interests(container, data):

    interests = data.get("interests", [])

    if not interests:
        return

    add_heading(container, "Interests")

    p = container.add_paragraph()

    p.space_before = Pt(2)
    p.space_after = Pt(2)

    if isinstance(interests, list):

        run = p.add_run(", ".join(interests))
        set_font(run, size=10)

    else:

        run = p.add_run(str(interests))
        set_font(run, size=10)