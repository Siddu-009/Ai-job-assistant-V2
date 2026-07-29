from docx.shared import Pt

from ..helpers import add_heading, set_font


def render_certifications(container, data):

    certifications = data.get("certifications", [])

    if not certifications:
        return

    add_heading(container, "Certifications")

    for cert in certifications:

        p = container.add_paragraph()

        p.space_before = Pt(2)
        p.space_after = Pt(2)

        if isinstance(cert, str):

            run = p.add_run("• " + cert)
            set_font(run, size=10)

            continue

        name = cert.get("name", "")
        issuer = cert.get("issuer", "")
        year = cert.get("year", "")

        title = p.add_run("• " + name)
        set_font(title, size=10, bold=True)

        details = []

        if issuer:
            details.append(issuer)

        if year:
            details.append(year)

        if details:

            extra = p.add_run(" | " + " | ".join(details))
            set_font(extra, size=10)