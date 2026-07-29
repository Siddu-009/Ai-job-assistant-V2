from docx.shared import Pt

from ..helpers import add_heading, set_font


def render_languages(container, data):

    languages = data.get("languages", [])

    if not languages:
        return

    add_heading(container, "Languages")

    p = container.add_paragraph()

    p.space_before = Pt(2)
    p.space_after = Pt(2)

    if isinstance(languages, list):

        run = p.add_run(", ".join(languages))
        set_font(run, size=10)

    else:

        run = p.add_run(str(languages))
        set_font(run, size=10)