from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt

from ..helpers import add_heading, set_font


def render_education(container, data):

    education = data.get("education", [])

    if not education:
        return

    add_heading(container, "Education")

    for edu in education:

        degree = (
            edu.get("degree")
            or edu.get("name")
            or ""
        )

        institution = (
            edu.get("institution")
            or edu.get("college")
            or edu.get("university")
            or ""
        )

        year = edu.get("year", "")
        cgpa = edu.get("cgpa", "")

        # ----------------------------------------
        # Degree + Year
        # ----------------------------------------

        p = container.add_paragraph()

        p.space_before = Pt(8)
        p.space_after = Pt(1)

        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(
            Inches(6.2),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.SPACES
        )

        left = p.add_run(degree)
        set_font(left, size=11, bold=True)

        if year:
            right = p.add_run("\t" + year)
            set_font(right, size=10)

        # ----------------------------------------
        # Institution
        # ----------------------------------------

        if institution:

            p = container.add_paragraph()

            p.space_before = Pt(0)
            p.space_after = Pt(1)

            run = p.add_run(institution)
            set_font(run, size=10)

        # ----------------------------------------
        # CGPA / Percentage
        # ----------------------------------------

        if cgpa:

            p = container.add_paragraph()

            p.space_before = Pt(0)
            p.space_after = Pt(2)

            label = p.add_run("CGPA / Percentage: ")
            set_font(label, size=10, bold=True)

            value = p.add_run(str(cgpa))
            set_font(value, size=10)

        container.add_paragraph()