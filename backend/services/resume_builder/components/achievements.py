from docx.shared import Pt

from ..helpers import add_heading, set_font


def render_achievements(container, data):

    achievements = data.get("achievements", [])

    if not achievements:
        return

    add_heading(container, "Achievements")

    for achievement in achievements:

        p = container.add_paragraph()

        p.space_before = Pt(1)
        p.space_after = Pt(1)

        run = p.add_run("• " + str(achievement))
        set_font(run, size=10)