from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from ..components.projects import render_projects
from ..components.education import render_education
from ..components.certifications import render_certifications
from ..components.achievements import render_achievements
from ..components.languages import render_languages
from ..components.interests import render_interests

from ..helpers import (
    set_font,
    add_heading,
    add_text,
    add_bullet,
    add_separator,
)

from ..styles import (
    PRIMARY,
    TITLE_SIZE,
    SUBTITLE_SIZE,
)


def build_classic_resume(data):

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(40)
    section.right_margin = Pt(40)

    # --------------------------------------------------
    # NAME
    # --------------------------------------------------

    personal = data.get("personal_info", {})

    name = (
        personal.get("name")
        or data.get("name")
        or ""
    )

    title = (
        personal.get("title")
        or data.get("title")
        or ""
    )

    email = (
        personal.get("email")
        or data.get("email")
        or ""
    )

    phone = (
        personal.get("phone")
        or data.get("phone")
        or ""
    )

    linkedin = (
        personal.get("linkedin")
        or data.get("linkedin")
        or ""
    )

    github = (
        personal.get("github")
        or data.get("github")
        or ""
    )

    portfolio = (
        personal.get("portfolio")
        or data.get("portfolio")
        or ""
    )

    location = (
        personal.get("location")
        or data.get("location")
        or ""
    )

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(name)

    set_font(
        run,
        size=TITLE_SIZE,
        bold=True,
        color=PRIMARY,
    )

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    if title:

        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(title)

        set_font(
            run,
            size=SUBTITLE_SIZE,
        )

    # --------------------------------------------------
    # CONTACT
    # --------------------------------------------------

    contact_items = []

    if phone:
        contact_items.append(phone)

    if email:
        contact_items.append(email)

    if location:
        contact_items.append(location)

    if linkedin:
        contact_items.append(linkedin)

    if github:
        contact_items.append(github)

    if portfolio:
        contact_items.append(portfolio)

    if contact_items:

        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(" | ".join(contact_items))

        set_font(run)

    add_separator(doc)

    # --------------------------------------------------
    # SUMMARY
    # --------------------------------------------------

    from ..components.summary import render_summary

    render_summary(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # SKILLS
    # --------------------------------------------------

    from ..components.skills import render_skills

    render_skills(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # EXPERIENCE
    # --------------------------------------------------

    from ..components.experience import render_experience

    render_experience(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # PROJECTS
    # --------------------------------------------------

    render_projects(doc, data)
    add_separator(doc)
    # --------------------------------------------------
    # EDUCATION
    # --------------------------------------------------

    render_education(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # CERTIFICATIONS
    # --------------------------------------------------

    render_certifications(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # ACHIEVEMENTS
    # --------------------------------------------------

    render_achievements(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # LANGUAGES
    # --------------------------------------------------

    render_languages(doc, data)
    add_separator(doc)

    # --------------------------------------------------
    # INTERESTS
    # --------------------------------------------------

    render_interests(doc, data)

    return doc