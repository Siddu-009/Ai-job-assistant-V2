from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from ..helpers import (
    set_font,
    add_heading,
    add_text,
    add_bullet,
)

from ..styles import (
    PRIMARY,
    TITLE_SIZE,
    SUBTITLE_SIZE,
)


def build_modern_resume(data):

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(36)
    section.right_margin = Pt(36)

    # ----------------------------------------
    # PERSONAL INFORMATION
    # ----------------------------------------

    personal = data.get("personal_info", {})

    name = (
        personal.get("name")
        or data.get("name")
        or ""
    )

    title = (
        personal.get("title")
        or data.get("title")
        or ""
    )

    email = (
        personal.get("email")
        or data.get("email")
        or ""
    )

    phone = (
        personal.get("phone")
        or data.get("phone")
        or ""
    )

    linkedin = (
        personal.get("linkedin")
        or data.get("linkedin")
        or ""
    )

    github = (
        personal.get("github")
        or data.get("github")
        or ""
    )

    portfolio = (
        personal.get("portfolio")
        or data.get("portfolio")
        or ""
    )

    location = (
        personal.get("location")
        or data.get("location")
        or ""
    )

    # ----------------------------------------
    # HEADER
    # ----------------------------------------

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(name)

    set_font(
        run,
        size=TITLE_SIZE,
        bold=True,
        color=PRIMARY
    )

    if title:

        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(title)

        set_font(
            run,
            size=SUBTITLE_SIZE
        )

    contact = []

    if phone:
        contact.append(phone)

    if email:
        contact.append(email)

    if location:
        contact.append(location)

    if linkedin:
        contact.append(linkedin)

    if github:
        contact.append(github)

    if portfolio:
        contact.append(portfolio)

    if contact:

        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(" | ".join(contact))

        set_font(run)

    # ----------------------------------------
    # MAIN TABLE
    # ----------------------------------------

    table = doc.add_table(
        rows=1,
        cols=2
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    left = table.cell(0, 0)

    right = table.cell(0, 1)

    # ========================================
    # LEFT COLUMN
    # ========================================

    add_heading(
        left,
        "Contact"
    )

    for item in contact:

        add_text(
            left,
            item
        )

    # ----------------------------------------
    # SKILLS
    # ----------------------------------------

    add_heading(
        left,
        "Technical Skills"
    )

    skill_categories = data.get(
        "skill_categories"
    )

    if skill_categories:

        for category in skill_categories:

            category_name = category.get(
                "category",
                ""
            )

            if category_name:

                add_text(
                    left,
                    category_name,
                    bold=True
                )

            for skill in category.get(
                "skills",
                []
            ):

                add_bullet(
                    left,
                    skill
                )

    else:

        skills = data.get(
            "skills",
            []
        )

        if isinstance(
            skills,
            list
        ):

            for skill in skills:

                add_bullet(
                    left,
                    skill
                )

        elif isinstance(
            skills,
            str
        ):

            for skill in skills.split(","):

                add_bullet(
                    left,
                    skill.strip()
                )

    # ----------------------------------------
    # CERTIFICATIONS
    # ----------------------------------------

    certifications = data.get(
        "certifications",
        []
    )

    if certifications:

        add_heading(
            left,
            "Certifications"
        )

        if isinstance(
            certifications,
            list
        ):

            for cert in certifications:

                if isinstance(cert, dict):

                    text = cert.get(
                        "name",
                        ""
                    )

                    issuer = cert.get(
                        "issuer",
                        ""
                    )

                    if issuer:
                        text += f" - {issuer}"

                    add_bullet(
                        left,
                        text
                    )

                else:

                    add_bullet(
                        left,
                        cert
                    )

        else:

            add_bullet(
                left,
                certifications
            )

    # ----------------------------------------
    # LANGUAGES
    # ----------------------------------------

    languages = data.get(
        "languages",
        []
    )

    if languages:

        add_heading(
            left,
            "Languages"
        )

        for language in languages:

            add_bullet(
                left,
                language
            )

    # ========================================
    # RIGHT COLUMN
    # ========================================

    summary = data.get(
        "summary",
        ""
    )

    if summary:

        add_heading(
            right,
            "Professional Summary"
        )

        add_text(
            right,
            summary
        )

    # ----------------------------------------
    # EXPERIENCE
    # ----------------------------------------

    experience = data.get("experience", [])

    if experience:

        add_heading(
            right,
            "Professional Experience"
        )

        if isinstance(experience, str):

            add_text(
                right,
                experience
            )

        else:

            for exp in experience:

                company = (
                    exp.get("company")
                    or exp.get("organization")
                    or exp.get("employer")
                    or exp.get("name")
                    or ""
                )

                designation = (
                    exp.get("designation")
                    or exp.get("title")
                    or exp.get("role")
                    or ""
                )

                location = exp.get(
                    "location",
                    ""
                )

                start_date = exp.get(
                    "start_date",
                    ""
                )

                end_date = exp.get(
                    "end_date",
                    ""
                )

                if company:

                    add_text(
                        right,
                        company,
                        bold=True
                    )

                if designation:

                    add_text(
                        right,
                        designation
                    )

                duration = " - ".join(
                    [
                        x
                        for x in [
                            start_date,
                            end_date
                        ]
                        if x
                    ]
                )

                if duration:

                    add_text(
                        right,
                        duration
                    )

                if location:

                    add_text(
                        right,
                        location
                    )

                description = exp.get(
                    "description",
                    []
                )

                if isinstance(
                    description,
                    list
                ):

                    for point in description:

                        add_bullet(
                            right,
                            point
                        )

                elif description:

                    add_bullet(
                        right,
                        description
                    )

                right.add_paragraph()

    # ----------------------------------------
    # PROJECTS
    # ----------------------------------------

    projects = data.get(
        "projects",
        []
    )

    if projects:

        add_heading(
            right,
            "Projects"
        )

        for project in projects:

            project_name = (
                project.get("name")
                or project.get("title")
                or ""
            )

            if project_name:

                add_text(
                    right,
                    project_name,
                    bold=True
                )

            role = project.get(
                "role",
                ""
            )

            if role:

                add_text(
                    right,
                    role
                )

            tech = (
                project.get("tech_stack")
                or project.get("technologies")
                or []
            )

            if tech:

                add_text(
                    right,
                    "Technologies: "
                    + ", ".join(tech)
                )

            description = project.get(
                "description",
                []
            )

            if isinstance(
                description,
                list
            ):

                for point in description:

                    add_bullet(
                        right,
                        point
                    )

            elif description:

                add_bullet(
                    right,
                    description
                    )

            right.add_paragraph()

    # ----------------------------------------
    # EDUCATION
    # ----------------------------------------

    education = data.get(
        "education",
        []
    )

    if education:

        add_heading(
            right,
            "Education"
        )

        for edu in education:

            degree = (
                edu.get("degree")
                or edu.get("name")
                or ""
            )

            institution = (
                edu.get("institution")
                or edu.get("college")
                or edu.get("university")
                or ""
            )

            if degree:

                add_text(
                    right,
                    degree,
                    bold=True
                )

            if institution:

                add_text(
                    right,
                    institution
                )

            cgpa = edu.get(
                "cgpa",
                ""
            )

            if cgpa:

                add_text(
                    right,
                    f"CGPA: {cgpa}"
                )

            year = edu.get(
                "year",
                ""
            )

            if year:

                add_text(
                    right,
                    f"Year: {year}"
                )

            right.add_paragraph()

    # ----------------------------------------
    # ACHIEVEMENTS
    # ----------------------------------------

    achievements = data.get(
        "achievements",
        []
    )

    if achievements:

        add_heading(
            right,
            "Achievements"
        )

        for achievement in achievements:

            add_bullet(
                right,
                achievement
            )

    # ----------------------------------------
    # INTERESTS
    # ----------------------------------------

    interests = data.get(
        "interests",
        []
    )

    if interests:

        add_heading(
            right,
            "Interests"
        )

        for interest in interests:

            add_bullet(
                right,
                interest
            )

    return doc