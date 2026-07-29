from docx.shared import Pt
from docx.oxml.ns import qn
from .styles import FONT_NAME, PRIMARY


def set_font(run, size=10, bold=False, color=None):

    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold

    if color:
        run.font.color.rgb = color


def add_heading(container, text):

    p = container.add_paragraph()

    p.space_before = Pt(8)
    p.space_after = Pt(4)

    run = p.add_run(text.upper())

    set_font(
        run,
        size=13,
        bold=True,
        color=PRIMARY
    )


def add_text(container, text, bold=False):

    if not text:
        return

    p = container.add_paragraph()

    p.space_after = Pt(2)

    run = p.add_run(str(text))

    set_font(
        run,
        size=10,
        bold=bold
    )


def add_bullet(container, text):

    if not text:
        return

    p = container.add_paragraph(
        style="List Bullet"
    )

    run = p.add_run(str(text))

    set_font(
        run,
        size=10
    )


def add_separator(container):

    p = container.add_paragraph()

    run = p.add_run("─" * 80)

    set_font(run, size=8)