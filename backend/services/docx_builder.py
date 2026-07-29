from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn


PRIMARY = RGBColor(33, 97, 140)


def set_font(run, size=10, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold

    if color:
        run.font.color.rgb = color


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    p.space_after = Pt(4)

    run = p.add_run(text.upper())

    set_font(
        run,
        size=13,
        bold=True,
        color=PRIMARY
    )


def add_line(doc):
    p = doc.add_paragraph()
    p.add_run("―" * 90)


def add_text(doc, text, bold=False):
    p = doc.add_paragraph()
    p.space_after = Pt(2)

    run = p.add_run(text)

    set_font(run, 10, bold)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")

    run = p.add_run(text)

    set_font(run, 10)


def build_resume(data):

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(40)
    section.right_margin = Pt(40)

    # ---------------- NAME ----------------

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(data.get("name", ""))

    set_font(
        run,
        size=22,
        bold=True,
        color=PRIMARY
    )

    # ---------------- TITLE ----------------

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(data.get("title", ""))

    set_font(run, 12)

    # ---------------- CONTACT ----------------

    p = doc.add_paragraph()

    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = p.add_run(
        f"{data.get('email','')} | "
        f"{data.get('phone','')} | "
        f"{data.get('linkedin','')} | "
        f"{data.get('github','')}"
    )

    set_font(run, 10)

    add_line(doc)

    # SUMMARY

    add_heading(doc, "Professional Summary")

    add_text(doc, data.get("summary", ""))

    add_line(doc)

    # SKILLS

    add_heading(doc, "Technical Skills")

    for skill in data.get("skills", []):
        add_bullet(doc, skill)

    add_line(doc)

    # EXPERIENCE

    add_heading(doc, "Professional Experience")

    experience = data.get("experience", [])

    if isinstance(experience, str):

        add_text(doc, experience)

    else:

        for exp in experience:

            add_text(
                doc,
                exp.get("name", ""),
                bold=True
            )

            if exp.get("title"):
                add_text(doc, exp["title"])

            desc = exp.get("description")

            if isinstance(desc, list):

                for d in desc:
                    add_bullet(doc, d)

            elif desc:
                add_bullet(doc, desc)

            doc.add_paragraph()

    add_line(doc)

    # PROJECTS

    add_heading(doc, "Projects")

    for project in data.get("projects", []):

        add_text(
            doc,
            project.get("name", ""),
            bold=True
        )

        tech = project.get("tech_stack", [])

        if tech:

            add_text(
                doc,
                "Technologies: "
                + ", ".join(tech)
            )

        if project.get("description"):

            add_bullet(
                doc,
                project["description"]
            )

        doc.add_paragraph()

    add_line(doc)

    # EDUCATION

    add_heading(doc, "Education")

    for edu in data.get("education", []):

        add_text(
            doc,
            edu.get("name", ""),
            bold=True
        )

        add_text(
            doc,
            edu.get("institution", "")
        )

        if edu.get("cgpa"):
            add_text(
                doc,
                f"CGPA: {edu['cgpa']}"
            )

        if edu.get("year"):
            add_text(
                doc,
                f"Year: {edu['year']}"
            )

        doc.add_paragraph()

    add_line(doc)

    # CERTIFICATIONS

    add_heading(doc, "Certifications")

    certs = data.get("certifications", [])

    if isinstance(certs, str):
        add_bullet(doc, certs)
    else:
        for cert in certs:
            add_bullet(doc, cert)

    return doc