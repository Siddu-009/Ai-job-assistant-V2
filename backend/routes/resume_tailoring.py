from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from services.ai_service import ai_chat
from services.notification_service import create_notification
from services.token_service import decode_token


router = APIRouter()


class ResumeTailorRequest(BaseModel):
    token: str
    resume: str
    job_description: str


BASE_DIR = Path(__file__).resolve().parent.parent

GENERATED_DIR = BASE_DIR / "generated"

GENERATED_DIR.mkdir(
    parents=True,
    exist_ok=True
)

DOCX_PATH = GENERATED_DIR / "tailored_resume.docx"


def create_docx_resume(
    resume_text: str,
    output_path: Path
):

    document = Document()

    # Page margins
    section = document.sections[0]

    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(45)
    section.right_margin = Pt(45)

    # Default font
    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # Process AI-generated resume
    lines = resume_text.splitlines()

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # Remove markdown symbols
        clean_line = line

        clean_line = clean_line.replace(
            "**",
            ""
        )

        clean_line = clean_line.replace(
            "__",
            ""
        )

        # Heading detection
        heading_words = [
            "professional summary",
            "summary",
            "skills",
            "technical skills",
            "experience",
            "professional experience",
            "projects",
            "education",
            "certifications",
            "achievements",
            "languages",
            "interests"
        ]

        if clean_line.lower().rstrip(":") in heading_words:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)

            run = paragraph.add_run(
                clean_line.rstrip(":")
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(12)

            continue

        # Main name detection
        if (
            len(clean_line) < 60
            and clean_line.isupper()
            and len(clean_line.split()) <= 6
        ):

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_PARAGRAPH_ALIGNMENT.CENTER
            )

            run = paragraph.add_run(
                clean_line
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(18)

            continue

        # Bullet points
        if (
            clean_line.startswith("- ")
            or clean_line.startswith("* ")
            or clean_line.startswith("• ")
        ):

            bullet_text = clean_line[2:].strip()

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            run = paragraph.add_run(
                bullet_text
            )

            run.font.name = "Arial"
            run.font.size = Pt(10)

            continue

        # Normal paragraph
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(
            clean_line
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)

    document.save(
        str(output_path)
    )


@router.post("/")
def tailor_resume(
    req: ResumeTailorRequest
):

    payload = decode_token(
        req.token
    )

    if not payload:

        return {
            "success": False,
            "message": "Invalid Token"
        }

    user_id = payload["user_id"]

    if not req.resume.strip():

        return {
            "success": False,
            "message": "Resume cannot be empty."
        }

    if not req.job_description.strip():

        return {
            "success": False,
            "message": "Job description cannot be empty."
        }

    prompt = f"""
You are an expert ATS Resume Writer.

Your task is to rewrite the user's existing resume
specifically for the provided job description.

IMPORTANT RULES:

1. Do NOT invent companies.
2. Do NOT invent degrees.
3. Do NOT invent certifications.
4. Do NOT invent job titles.
5. Do NOT invent years of experience.
6. Do NOT invent technologies the user has never mentioned.
7. You may reorganize and improve the existing information.
8. You may naturally include relevant keywords from the job description
   only when they are supported by the user's existing experience or skills.
9. Use strong action verbs.
10. Make the resume ATS friendly.
11. Keep the resume professional and concise.
12. Do not include explanations before or after the resume.
13. Return ONLY the final tailored resume.

STRUCTURE:

FULL NAME
PROFESSIONAL TITLE
CONTACT INFORMATION

PROFESSIONAL SUMMARY

TECHNICAL SKILLS

PROFESSIONAL EXPERIENCE

PROJECTS

EDUCATION

CERTIFICATIONS

ACHIEVEMENTS

LANGUAGES

INTERESTS

==============================
ORIGINAL RESUME
==============================

{req.resume}

==============================
JOB DESCRIPTION
==============================

{req.job_description}

==============================
FINAL INSTRUCTION
==============================

Create the final ATS-optimized resume now.

Return ONLY the resume content.
"""


    try:

        result = ai_chat(
            prompt
        )

        if not result:

            raise Exception(
                "AI returned an empty response."
            )

        # Detect AI connection errors
        if result.startswith(
            "Unable to connect to Ollama"
        ):

            raise HTTPException(
                status_code=503,
                detail=result
            )

        if result.startswith(
            "The AI model took too long"
        ):

            raise HTTPException(
                status_code=504,
                detail=result
            )

        # Create DOCX
        create_docx_resume(
            result,
            DOCX_PATH
        )

        if not DOCX_PATH.exists():

            raise Exception(
                "Tailored DOCX was not created."
            )

        create_notification(
            user_id=user_id,
            title="📄 Resume Tailored",
            message=(
                "Your resume has been tailored "
                "successfully."
            ),
            notification_type="resume",
            link="/resume-tailoring"
        )

        return {
            "success": True,
            "message": (
                "Tailored resume generated "
                "successfully."
            ),
            "tailored_resume": result,
            "download_url": (
                "/api/resume-tailoring/download"
            )
        }

    except HTTPException:
        raise

    except Exception as e:

        print(
            "Resume Tailoring Error:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


@router.get("/download")
def download_tailored_resume():

    if not DOCX_PATH.exists():

        raise HTTPException(
            status_code=404,
            detail=(
                "Tailored resume has not "
                "been generated yet."
            )
        )

    return FileResponse(
        path=str(DOCX_PATH),
        filename="tailored_resume.docx",
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )